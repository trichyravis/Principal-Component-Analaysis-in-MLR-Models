
"""
THE MOUNTAIN PATH - WORLD OF FINANCE
Principal Component Analysis (PCA) - Interactive Learning App
Prof. V. Ravichandran | mountainpathacademy.com

Multi-tab Streamlit application for learning PCA from basics to case studies.
Run: streamlit run pca_learning_app.py
"""

import streamlit as st
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
from plotly.subplots import make_subplots
import io

# =============================================================================
# PAGE CONFIG & MOUNTAIN PATH DESIGN
# =============================================================================
st.set_page_config(
    page_title="PCA Learning Guide | The Mountain Path",
    page_icon="🏔️",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Mountain Path Academy Color Scheme
DARK_BLUE = "#003366"
LIGHT_BLUE = "#ADD8E6"
GOLD = "#FFD700"
WHITE = "#FFFFFF"
DARK_BG = "#F8FAFC"
CARD_BG = "#FFFFFF"

# Custom CSS for Mountain Path design
st.markdown(f"""
<style>
    /* Main background */
    .stApp {{
        background-color: {DARK_BG};
    }}

    /* Sidebar styling */
    section[data-testid="stSidebar"] {{
        background-color: {DARK_BLUE};
    }}
    section[data-testid="stSidebar"] .stMarkdown p,
    section[data-testid="stSidebar"] .stMarkdown li,
    section[data-testid="stSidebar"] .stMarkdown h1,
    section[data-testid="stSidebar"] .stMarkdown h2,
    section[data-testid="stSidebar"] .stMarkdown h3 {{
        color: {WHITE};
    }}
    section[data-testid="stSidebar"] .stRadio label p {{
        color: {WHITE};
        font-size: 1rem;
    }}
    section[data-testid="stSidebar"] .stRadio div[role="radiogroup"] label:hover {{
        background-color: rgba(255,215,0,0.15);
        border-radius: 5px;
    }}

    /* Header banner */
    .mountain-header {{
        background: linear-gradient(135deg, {DARK_BLUE}, #004488);
        color: white;
        padding: 1.5rem 2rem;
        border-radius: 10px;
        margin-bottom: 1.5rem;
        box-shadow: 0 4px 15px rgba(0,51,102,0.3);
    }}
    .mountain-header h1 {{
        color: white !important;
        margin: 0;
        font-size: 1.8rem;
    }}
    .mountain-header p {{
        color: {LIGHT_BLUE};
        margin: 0.3rem 0 0 0;
        font-size: 1rem;
    }}

    /* Definition box */
    .def-box {{
        background-color: #f0f8ff;
        border-left: 5px solid {DARK_BLUE};
        padding: 1rem 1.2rem;
        border-radius: 0 8px 8px 0;
        margin: 1rem 0;
    }}
    .def-box h4 {{
        color: {DARK_BLUE};
        margin-top: 0;
    }}

    /* Gold highlight box */
    .gold-box {{
        background: linear-gradient(135deg, #FFFDE7, #FFF9C4);
        border-left: 5px solid {GOLD};
        padding: 1rem 1.2rem;
        border-radius: 0 8px 8px 0;
        margin: 1rem 0;
    }}
    .gold-box h4 {{
        color: #7B6B00;
        margin-top: 0;
    }}

    /* Insight box */
    .insight-box {{
        background-color: #E8F0FE;
        border: 2px solid {DARK_BLUE};
        padding: 1rem 1.2rem;
        border-radius: 8px;
        margin: 1rem 0;
    }}
    .insight-box h4 {{
        color: {DARK_BLUE};
        margin-top: 0;
    }}

    /* Warning box */
    .warn-box {{
        background-color: #FFF3E0;
        border-left: 5px solid #E65100;
        padding: 1rem 1.2rem;
        border-radius: 0 8px 8px 0;
        margin: 1rem 0;
    }}

    /* Calculation box */
    .calc-box {{
        background-color: #E8F5E9;
        border-left: 5px solid #2E7D32;
        padding: 1rem 1.2rem;
        border-radius: 0 8px 8px 0;
        margin: 1rem 0;
        font-family: 'Courier New', monospace;
    }}

    /* Tab styling */
    .stTabs [data-baseweb="tab-list"] {{
        gap: 4px;
        background-color: {DARK_BLUE};
        border-radius: 10px;
        padding: 5px;
    }}
    .stTabs [data-baseweb="tab"] {{
        color: white;
        background-color: transparent;
        border-radius: 8px;
        padding: 10px 20px;
        font-weight: 600;
    }}
    .stTabs [aria-selected="true"] {{
        background-color: {GOLD} !important;
        color: {DARK_BLUE} !important;
    }}
    .stTabs [data-baseweb="tab"]:hover {{
        background-color: rgba(255,215,0,0.3);
    }}
    .stTabs [data-baseweb="tab-panel"] {{
        padding-top: 1.5rem;
    }}

    /* Metric cards */
    .metric-card {{
        background: white;
        border: 2px solid {DARK_BLUE};
        border-radius: 10px;
        padding: 1rem;
        text-align: center;
        box-shadow: 0 2px 8px rgba(0,0,0,0.08);
    }}
    .metric-card h2 {{
        color: {DARK_BLUE};
        margin: 0;
        font-size: 2rem;
    }}
    .metric-card p {{
        color: #555;
        margin: 0.3rem 0 0 0;
        font-size: 0.9rem;
    }}

    /* Footer */
    .footer {{
        background-color: {DARK_BLUE};
        color: white;
        padding: 1rem;
        border-radius: 10px;
        text-align: center;
        margin-top: 2rem;
    }}

    /* Hide default streamlit elements for cleaner look */
    #MainMenu {{visibility: hidden;}}
    header {{visibility: hidden;}}
</style>
""", unsafe_allow_html=True)


# =============================================================================
# HELPER FUNCTIONS
# =============================================================================
def mountain_header(title, subtitle=""):
    st.markdown(f"""
    <div class="mountain-header">
        <h1>🏔️ {title}</h1>
        <p>{subtitle}</p>
    </div>
    """, unsafe_allow_html=True)

def def_box(title, content):
    st.markdown(f"""
    <div class="def-box">
        <h4>📘 {title}</h4>
        <p>{content}</p>
    </div>
    """, unsafe_allow_html=True)

def gold_box(title, content):
    st.markdown(f"""
    <div class="gold-box">
        <h4>💡 {title}</h4>
        <p>{content}</p>
    </div>
    """, unsafe_allow_html=True)

def insight_box(title, content):
    st.markdown(f"""
    <div class="insight-box">
        <h4>🔍 {title}</h4>
        <p>{content}</p>
    </div>
    """, unsafe_allow_html=True)

def warn_box(content):
    st.markdown(f"""
    <div class="warn-box">
        <strong>⚠️ Warning:</strong> {content}
    </div>
    """, unsafe_allow_html=True)

def calc_box(title, content):
    st.markdown(f"""
    <div class="calc-box">
        <h4>🧮 {title}</h4>
        {content}
    </div>
    """, unsafe_allow_html=True)

def metric_card(value, label):
    st.markdown(f"""
    <div class="metric-card">
        <h2>{value}</h2>
        <p>{label}</p>
    </div>
    """, unsafe_allow_html=True)

def footer():
    st.markdown("""
    <div class="footer">
        <strong>The Mountain Path — World of Finance</strong><br>
        Prof. V. Ravichandran | mountainpathacademy.com<br>
        <em>Bridging Theory with Practice • Excellence in Financial Education</em>
    </div>
    """, unsafe_allow_html=True)


# =============================================================================
# SIDEBAR
# =============================================================================
with st.sidebar:
    st.markdown("""
    ## 🏔️ THE MOUNTAIN PATH
    ### World of Finance
    ---
    """)
    st.markdown("### 📚 PCA Learning Guide")
    st.markdown("Navigate through the tabs to explore PCA from basics to advanced case studies.")
    st.markdown("---")
    st.markdown("""
    **Prof. V. Ravichandran**
    - 28+ Years Corporate Finance
    - 10+ Years Academic Excellence
    """)
    st.markdown("---")
    st.markdown("🌐 [mountainpathacademy.com](https://mountainpathacademy.com)")


# =============================================================================
# MAIN TABS
# =============================================================================
mountain_header(
    "Principal Component Analysis (PCA)",
    "A Complete Interactive Learning Guide — From Intuition to Implementation"
)

tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
    "📖 Basics",
    "🧠 Concepts & Math",
    "📊 Interactive Illustrations",
    "🏠 Case Study",
    "🧪 Try It Yourself",
    "🔬 Lab Exercise"
])


# =============================================================================
# TAB 1: BASICS
# =============================================================================
with tab1:
    st.markdown(f"## Why PCA Matters")

    col1, col2, col3 = st.columns(3)
    with col1:
        metric_card("96.3%", "Variance captured by just PC1 in our 2-var example")
    with col2:
        metric_card("99.2%", "Variance captured by 3 of 5 PCs in the case study")
    with col3:
        metric_card("0.000", "Correlation between PCs (by construction)")

    st.markdown("---")

    # MLR Recap
    st.markdown("### 1. Multiple Linear Regression (MLR) — The Starting Point")

    def_box("What is MLR?",
        "MLR models the relationship between <strong>one</strong> dependent variable (Y) and "
        "<strong>two or more</strong> independent variables (X₁, X₂, ..., Xₚ). "
        "It finds the best-fit hyperplane through the data by minimising squared errors.")

    st.latex(r"Y = \beta_0 + \beta_1 X_1 + \beta_2 X_2 + \cdots + \beta_p X_p + \varepsilon")

    with st.expander("📋 Click to see all 5 MLR assumptions"):
        assumptions = pd.DataFrame({
            "Assumption": ["Linearity", "Independence", "Homoscedasticity", "Normality of Errors",
                           "No Multicollinearity ⭐"],
            "What It Means": [
                "Y vs X relationship is a straight line",
                "Observations don't influence each other",
                "Error spread is constant across all X values",
                "Residuals follow a bell curve",
                "Predictors are NOT highly correlated with each other"
            ],
            "PCA Relevant?": ["No", "No", "No", "No", "YES — This is what PCA fixes!"]
        })
        st.dataframe(assumptions, use_container_width=True, hide_index=True)

    st.markdown("---")

    # The Problem
    st.markdown("### 2. The Problem — Multicollinearity")

    def_box("What is Multicollinearity?",
        "When two or more predictor variables are <strong>highly correlated</strong>, they carry "
        "overlapping information. OLS cannot decide which variable deserves credit for predicting Y.")

    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**Consequences:**")
        st.markdown("""
        - 📈 Coefficient estimates swing wildly
        - 📊 Standard errors inflate → predictors look insignificant
        - 🔄 Signs of coefficients may flip (nonsensical)
        - 💥 (X'X) matrix nearly singular → inverse explodes
        - 🎭 High R² hides unreliable individual coefficients
        """)
    with col2:
        st.markdown("**Detection Methods:**")
        st.markdown("""
        - Correlation matrix: |r| > 0.8 → suspect
        - **VIF > 10** → serious problem
        - Eigenvalues near zero → near-singular
        - Condition number > 30 → trouble
        """)
        st.latex(r"\text{VIF}_j = \frac{1}{1 - R_j^2}")

    gold_box("The Fix",
        "PCA transforms correlated predictors into <strong>new, uncorrelated</strong> variables "
        "(principal components). These can safely be used in MLR — this combination is called "
        "<strong>Principal Component Regression (PCR)</strong>.")

    st.markdown("---")

    # What is PCA
    st.markdown("### 3. What Is PCA? — The One-Sentence Answer")

    insight_box("PCA in One Sentence",
        "PCA <strong>rotates</strong> your coordinate axes to align with the directions of "
        "<strong>maximum data spread</strong> (variance), producing new uncorrelated variables "
        "ordered from most to least informative.")

    st.markdown("**The Three Magic Properties of Principal Components:**")
    prop_cols = st.columns(3)
    with prop_cols[0]:
        st.markdown(f"""
        <div class="metric-card">
            <h2 style="font-size:2.5rem;">⊥</h2>
            <h3 style="color:{DARK_BLUE};">Orthogonal</h3>
            <p>PCs are perpendicular → zero correlation among them</p>
        </div>
        """, unsafe_allow_html=True)
    with prop_cols[1]:
        st.markdown(f"""
        <div class="metric-card">
            <h2 style="font-size:2.5rem;">📊</h2>
            <h3 style="color:{DARK_BLUE};">Ordered</h3>
            <p>PC1 explains most variance, PC2 next, and so on</p>
        </div>
        """, unsafe_allow_html=True)
    with prop_cols[2]:
        st.markdown(f"""
        <div class="metric-card">
            <h2 style="font-size:2.5rem;">🎯</h2>
            <h3 style="color:{DARK_BLUE};">Max Variance</h3>
            <p>Each PC captures as much remaining variance as possible</p>
        </div>
        """, unsafe_allow_html=True)

    footer()


# =============================================================================
# TAB 2: CONCEPTS & MATH
# =============================================================================
with tab2:
    st.markdown("## The Mathematics Behind PCA — Step by Step")

    concept_tab = st.radio(
        "Select a concept to explore:",
        ["Variance & Covariance", "The Covariance Matrix", "Eigenvalues & Eigenvectors",
         "Eigenvalue Calculation (2×2)", "Explained Variance", "PC Scores", "The 7-Step Recipe"],
        horizontal=True
    )

    if concept_tab == "Variance & Covariance":
        st.markdown("### Variance — Measuring Spread")
        def_box("Variance",
            "Variance measures how far data points are spread out from their mean:<br>"
            "<code>Var(X) = (1/(n-1)) × Σ(Xᵢ - X̄)²</code>")

        st.latex(r"\text{Var}(X) = s^2 = \frac{1}{n-1}\sum_{i=1}^{n}(X_i - \bar{X})^2")

        st.markdown("### Covariance — How Two Variables Move Together")
        def_box("Covariance",
            "Covariance measures whether two variables tend to increase together (positive), "
            "move in opposite directions (negative), or have no relationship (≈ 0):<br>"
            "<code>Cov(X, Z) = (1/(n-1)) × Σ(Xᵢ - X̄)(Zᵢ - Z̄)</code>")

        st.latex(r"\text{Cov}(X, Z) = \frac{1}{n-1}\sum_{i=1}^{n}(X_i - \bar{X})(Z_i - \bar{Z})")

        # Interactive covariance visualization
        st.markdown("#### Interactive: See How Covariance Works")
        corr_val = st.slider("Set correlation strength:", -1.0, 1.0, 0.8, 0.05)

        np.random.seed(42)
        n_pts = 100
        x = np.random.randn(n_pts)
        y = corr_val * x + np.sqrt(1 - corr_val**2) * np.random.randn(n_pts)

        fig = go.Figure()
        fig.add_trace(go.Scatter(x=x, y=y, mode='markers',
            marker=dict(color=DARK_BLUE, size=8, opacity=0.7),
            name='Data Points'))
        fig.add_hline(y=0, line_dash="dash", line_color="gray", opacity=0.5)
        fig.add_vline(x=0, line_dash="dash", line_color="gray", opacity=0.5)
        fig.update_layout(
            title=f"Scatter Plot — Correlation r = {corr_val:.2f}, Cov = {np.cov(x,y)[0,1]:.3f}",
            xaxis_title="X", yaxis_title="Z",
            template="plotly_white", height=400,
            font=dict(family="Times New Roman"))
        st.plotly_chart(fig, use_container_width=True)

    elif concept_tab == "The Covariance Matrix":
        st.markdown("### The Covariance Matrix Σ")
        def_box("Covariance Matrix",
            "For p variables, the covariance matrix is a p×p symmetric matrix where:<br>"
            "• <strong>Diagonal entries</strong> = variance of each variable<br>"
            "• <strong>Off-diagonal entries</strong> = covariance between pairs")

        st.latex(r"""\boldsymbol{\Sigma} = \begin{pmatrix}
            \text{Var}(X_1) & \text{Cov}(X_1, X_2) \\
            \text{Cov}(X_2, X_1) & \text{Var}(X_2)
        \end{pmatrix}""")

        calc_box("Our 2-Variable Example",
            "<code>Σ = [[0.6166, 0.6154], [0.6154, 0.7166]]</code><br><br>"
            "• Var(X₁) = 0.6166 &nbsp;&nbsp; • Var(X₂) = 0.7166<br>"
            "• Cov(X₁,X₂) = 0.6154 (strongly positive → they move together)<br><br>"
            "<strong>PCA's goal:</strong> Find a rotation that <em>diagonalises</em> this matrix.")

    elif concept_tab == "Eigenvalues & Eigenvectors":
        st.markdown("### Eigenvalues & Eigenvectors — The Core of PCA")

        def_box("Definition",
            "For a square matrix <strong>A</strong>, if there exists a non-zero vector <strong>v</strong> "
            "and scalar λ such that <strong>Av = λv</strong>, then:<br>"
            "• <strong>v</strong> is an <em>eigenvector</em> (a direction the matrix doesn't rotate, only stretches)<br>"
            "• <strong>λ</strong> is the <em>eigenvalue</em> (the stretching factor)")

        st.latex(r"\mathbf{A} \cdot \mathbf{v} = \lambda \cdot \mathbf{v}")

        col1, col2 = st.columns(2)
        with col1:
            insight_box("In PCA Context",
                "<strong>Eigenvectors</strong> of the covariance matrix = directions of principal components (new axes).<br>"
                "<strong>Eigenvalues</strong> = how much variance (information) lies along each direction.<br>"
                "Largest λ → PC1 (most important). Smallest λ → last PC (least important, often noise).")
        with col2:
            # Eigenvector visualization
            fig = go.Figure()
            # General vector
            fig.add_trace(go.Scatter(x=[0, 2], y=[0, 1], mode='lines+markers',
                line=dict(color=DARK_BLUE, width=3), name='Original v',
                marker=dict(size=[0, 10], symbol='arrow-bar-up')))
            fig.add_trace(go.Scatter(x=[0, 4], y=[0, 2], mode='lines+markers',
                line=dict(color='red', width=3, dash='dash'), name='Av = λv (stretched, same direction)',
                marker=dict(size=[0, 10])))
            fig.update_layout(
                title="Eigenvector: A only stretches, doesn't rotate",
                xaxis_title="", yaxis_title="",
                template="plotly_white", height=300,
                showlegend=True)
            st.plotly_chart(fig, use_container_width=True)

    elif concept_tab == "Eigenvalue Calculation (2×2)":
        st.markdown("### Complete Eigenvalue Calculation for a 2×2 Matrix")

        st.markdown("For our covariance matrix:")
        st.latex(r"""\boldsymbol{\Sigma} = \begin{pmatrix} 0.6166 & 0.6154 \\
            0.6154 & 0.7166 \end{pmatrix} \quad\text{where } a=0.6166,\ b=0.6154,\ c=0.6154,\ d=0.7166""")

        st.markdown("#### The Closed-Form Formula")
        st.latex(r"\lambda^2 - \text{Trace} \cdot \lambda + \text{Det} = 0")
        st.latex(r"\lambda = \frac{\text{Trace} \pm \sqrt{\text{Trace}^2 - 4 \cdot \text{Det}}}{2}")

        st.markdown("#### Step-by-Step Calculation")

        steps = {
            "Step 1: Trace (a + d)": {
                "formula": "Trace = a + d = 0.6166 + 0.7166",
                "result": 1.3332,
                "meaning": "Sum of diagonal = total variance across both variables"
            },
            "Step 2: Determinant (ad − bc)": {
                "formula": "Det = (0.6166)(0.7166) − (0.6154)(0.6154) = 0.4417 − 0.3787",
                "result": 0.0630,
                "meaning": "How 'non-singular' the matrix is. Small Det = high correlation"
            },
            "Step 3: Discriminant": {
                "formula": "Disc = Trace² − 4×Det = (1.3332)² − 4(0.0630) = 1.7774 − 0.2520",
                "result": 1.5254,
                "meaning": "The term under the square root in the quadratic formula"
            },
            "Step 4: √Discriminant": {
                "formula": "√1.5254",
                "result": 1.2351,
                "meaning": "Needed for the ± in the quadratic formula"
            },
            "Step 5: λ₁ (largest eigenvalue)": {
                "formula": "λ₁ = (Trace + √Disc) / 2 = (1.3332 + 1.2351) / 2",
                "result": 1.2841,
                "meaning": "Variance captured by PC1 — the dominant direction"
            },
            "Step 6: λ₂ (smallest eigenvalue)": {
                "formula": "λ₂ = (Trace − √Disc) / 2 = (1.3332 − 1.2351) / 2",
                "result": 0.0491,
                "meaning": "Variance captured by PC2 — the minor direction"
            },
        }

        for step_name, details in steps.items():
            with st.expander(f"**{step_name} = {details['result']:.4f}**", expanded=True):
                st.code(details["formula"])
                st.markdown(f"**Result: {details['result']:.4f}**")
                st.markdown(f"*{details['meaning']}*")

        gold_box("Verification",
            "λ₁ + λ₂ = 1.2841 + 0.0491 = <strong>1.3332 = Trace</strong> ✓<br>"
            "The sum of eigenvalues always equals the trace (total variance is preserved)!")

        # Summary table
        st.markdown("#### Complete Summary")
        summary_df = pd.DataFrame({
            "Quantity": ["Trace (a+d)", "Determinant (ad−bc)", "Discriminant",
                         "√Discriminant", "λ₁ (largest)", "λ₂ (smallest)"],
            "Formula": ["Sum of diagonal", "Product of diag − product of off-diag",
                         "Trace² − 4×Det", "", "(Trace + √Disc)/2", "(Trace − √Disc)/2"],
            "Value": [1.3332, 0.0630, 1.5254, 1.2351, 1.2841, 0.0491]
        })
        st.dataframe(summary_df, use_container_width=True, hide_index=True)

    elif concept_tab == "Explained Variance":
        st.markdown("### Explained Variance — How Important Is Each PC?")
        st.latex(r"\text{Explained Variance Ratio}_i = \frac{\lambda_i}{\lambda_1 + \lambda_2 + \cdots + \lambda_p}")

        calc_box("2-Variable Example",
            "Total variance = λ₁ + λ₂ = 1.2841 + 0.0491 = 1.3332<br><br>"
            "<strong>PC1% = 1.2841 / 1.3332 = 96.32%</strong><br>"
            "<strong>PC2% = 0.0491 / 1.3332 = 3.68%</strong><br><br>"
            "PC1 alone captures <strong>96.3%</strong> of all information!")

        # Interactive bar chart
        fig = go.Figure()
        fig.add_trace(go.Bar(
            x=["PC1", "PC2"], y=[96.32, 3.68],
            marker_color=[DARK_BLUE, LIGHT_BLUE],
            text=["96.32%", "3.68%"], textposition='outside',
            textfont=dict(size=16, color=DARK_BLUE, family="Times New Roman")))
        fig.update_layout(
            title="Explained Variance — 2-Variable Example",
            yaxis_title="Variance Explained (%)", yaxis_range=[0, 105],
            template="plotly_white", height=400,
            font=dict(family="Times New Roman"))
        st.plotly_chart(fig, use_container_width=True)

    elif concept_tab == "PC Scores":
        st.markdown("### PC Scores — Coordinates in the New System")
        def_box("What Are PC Scores?",
            "PC scores are the <strong>coordinates</strong> of each data point in the new PC "
            "coordinate system. Computed by projecting centred data onto eigenvectors:<br><br>"
            "<code>PC1 scoreᵢ = (X₁ᵢ − X̄₁) × v₁₁ + (X₂ᵢ − X̄₂) × v₁₂</code>")

        calc_box("Worked Calculation: Observation 1",
            "X₁ = 2.5, X₂ = 2.4 &nbsp;|&nbsp; Centred: 0.69, 0.49<br><br>"
            "PC1 = (0.69)(0.678) + (0.49)(0.735) = 0.468 + 0.360 = <strong>0.828</strong><br>"
            "PC2 = (0.69)(0.735) + (0.49)(−0.678) = 0.507 − 0.332 = <strong>0.175</strong>")

        insight_box("Sanity Check",
            "Correlation between all PC1 scores and all PC2 scores ≈ <strong>0.000</strong>. "
            "The new variables are <em>uncorrelated</em>, exactly as promised!")

    elif concept_tab == "The 7-Step Recipe":
        st.markdown("### The Complete PCA Workflow")

        steps_data = [
            ("1️⃣", "Organise Data", "Arrange as matrix X (n×p). Do NOT include Y."),
            ("2️⃣", "Standardise", "Z = (X − μ) / σ. Critical when variables have different scales!"),
            ("3️⃣", "Correlation Matrix", "Compute the p×p correlation matrix."),
            ("4️⃣", "Eigen-decompose", "Get eigenvalues (λ's) and eigenvectors (v's)."),
            ("5️⃣", "Choose k", "Cumulative variance ≥ 80–95%, Kaiser's rule, scree plot, or CV."),
            ("6️⃣", "Project Data", "Z_scores = Z_std × V_k → n×k score matrix."),
            ("7️⃣", "Use in MLR", "Regress Y on PC scores. Uncorrelated by construction!"),
        ]

        for emoji, title, desc in steps_data:
            st.markdown(f"""
            <div style="background: white; border-left: 4px solid {DARK_BLUE}; padding: 0.8rem 1rem;
                        margin: 0.5rem 0; border-radius: 0 8px 8px 0; box-shadow: 0 1px 4px rgba(0,0,0,0.06);">
                <strong style="color:{DARK_BLUE}; font-size:1.1rem;">{emoji} {title}</strong><br>
                <span style="color:#333;">{desc}</span>
            </div>
            """, unsafe_allow_html=True)

        warn_box("Common mistake: forgetting to standardise. A variable in thousands "
                 "will dominate PC1 purely because of scale, not because it's more important!")

    footer()


# =============================================================================
# TAB 3: INTERACTIVE ILLUSTRATIONS
# =============================================================================
with tab3:
    st.markdown("## Interactive Visualisations")

    viz_choice = st.radio(
        "Select visualisation:",
        ["Before & After PCA (2D)", "Scree Plot Explorer",
         "Loadings Heatmap", "Explained Variance Cumulative"],
        horizontal=True
    )

    if viz_choice == "Before & After PCA (2D)":
        st.markdown("### See PCA in Action — Rotate the Axes")

        # 2-variable data from the workbook
        X1 = np.array([2.5, 0.5, 2.2, 1.9, 3.1, 2.3, 2.0, 1.0, 1.5, 1.1])
        X2 = np.array([2.4, 0.7, 2.9, 2.2, 3.0, 2.7, 1.6, 1.1, 1.6, 0.9])
        X1c = X1 - X1.mean()
        X2c = X2 - X2.mean()

        cov_mat = np.cov(X1, X2)
        eigenvalues, eigenvectors = np.linalg.eig(cov_mat)
        idx = eigenvalues.argsort()[::-1]
        eigenvalues = eigenvalues[idx]
        eigenvectors = eigenvectors[:, idx]

        pc1_scores = X1c * eigenvectors[0, 0] + X2c * eigenvectors[1, 0]
        pc2_scores = X1c * eigenvectors[0, 1] + X2c * eigenvectors[1, 1]

        fig = make_subplots(rows=1, cols=2,
            subplot_titles=(
                f"Original Data (r = {np.corrcoef(X1, X2)[0,1]:.3f})",
                f"PC Scores (r = {np.corrcoef(pc1_scores, pc2_scores)[0,1]:.6f})"
            ), horizontal_spacing=0.12)

        # Original data
        fig.add_trace(go.Scatter(
            x=X1, y=X2, mode='markers',
            marker=dict(color=DARK_BLUE, size=12, line=dict(color='white', width=1.5)),
            name='Original', showlegend=False), row=1, col=1)

        # Add PC direction arrows on original plot
        center = [X1.mean(), X2.mean()]
        scale1 = np.sqrt(eigenvalues[0]) * 1.5
        scale2 = np.sqrt(eigenvalues[1]) * 1.5
        fig.add_trace(go.Scatter(
            x=[center[0] - eigenvectors[0,0]*scale1, center[0] + eigenvectors[0,0]*scale1],
            y=[center[1] - eigenvectors[1,0]*scale1, center[1] + eigenvectors[1,0]*scale1],
            mode='lines', line=dict(color='red', width=3, dash='dash'),
            name='PC1 direction', showlegend=True), row=1, col=1)
        fig.add_trace(go.Scatter(
            x=[center[0] - eigenvectors[0,1]*scale2, center[0] + eigenvectors[0,1]*scale2],
            y=[center[1] - eigenvectors[1,1]*scale2, center[1] + eigenvectors[1,1]*scale2],
            mode='lines', line=dict(color=GOLD, width=3, dash='dash'),
            name='PC2 direction', showlegend=True), row=1, col=1)

        # PC scores
        fig.add_trace(go.Scatter(
            x=pc1_scores, y=pc2_scores, mode='markers',
            marker=dict(color=DARK_BLUE, size=12, line=dict(color='white', width=1.5)),
            name='PC Scores', showlegend=False), row=1, col=2)

        fig.update_xaxes(title_text="X₁", row=1, col=1)
        fig.update_yaxes(title_text="X₂", row=1, col=1)
        fig.update_xaxes(title_text="PC1", row=1, col=2)
        fig.update_yaxes(title_text="PC2", row=1, col=2)
        fig.update_layout(
            template="plotly_white", height=500,
            font=dict(family="Times New Roman", size=13))
        st.plotly_chart(fig, use_container_width=True)

        gold_box("What You See",
            "Left: the diagonal stretch shows high correlation (r = 0.93). "
            "The red dashed line = PC1 direction (maximum spread). "
            "Right: after PCA, the cloud is horizontal along PC1 — <strong>correlation is gone</strong>.")

    elif viz_choice == "Scree Plot Explorer":
        st.markdown("### Scree Plot — How to Choose the Number of Components")

        dataset = st.selectbox("Select dataset:", ["2-Variable Example", "5-Variable Case Study"])

        if dataset == "2-Variable Example":
            evals = np.array([1.2841, 0.0491])
            labels = ["PC1", "PC2"]
        else:
            evals = np.array([4.602, 0.228, 0.132, 0.035, 0.002])
            labels = ["PC1", "PC2", "PC3", "PC4", "PC5"]

        var_pct = evals / evals.sum() * 100
        cum_pct = np.cumsum(var_pct)

        fig = make_subplots(specs=[[{"secondary_y": True}]])

        fig.add_trace(go.Bar(
            x=labels, y=var_pct, name="Individual %",
            marker_color=DARK_BLUE, opacity=0.8,
            text=[f"{v:.1f}%" for v in var_pct], textposition='outside'), secondary_y=False)

        fig.add_trace(go.Scatter(
            x=labels, y=cum_pct, name="Cumulative %",
            mode='lines+markers+text',
            line=dict(color='red', width=3),
            marker=dict(size=10, color='red'),
            text=[f"{v:.1f}%" for v in cum_pct], textposition='top center'), secondary_y=True)

        # Kaiser's rule line
        fig.add_hline(y=100/len(evals), line_dash="dot", line_color=GOLD,
            annotation_text=f"Kaiser threshold ({100/len(evals):.0f}%)", secondary_y=False)

        fig.update_yaxes(title_text="Individual Variance (%)", range=[0, 105], secondary_y=False)
        fig.update_yaxes(title_text="Cumulative Variance (%)", range=[0, 105], secondary_y=True)
        fig.update_layout(
            title=f"Scree Plot — {dataset}",
            template="plotly_white", height=500,
            font=dict(family="Times New Roman", size=13),
            legend=dict(orientation="h", yanchor="bottom", y=1.02))
        st.plotly_chart(fig, use_container_width=True)

        # Eigenvalue table
        ev_df = pd.DataFrame({
            "PC": labels,
            "Eigenvalue (λ)": [f"{e:.4f}" for e in evals],
            "Variance %": [f"{v:.2f}%" for v in var_pct],
            "Cumulative %": [f"{c:.2f}%" for c in cum_pct],
        })
        st.dataframe(ev_df, use_container_width=True, hide_index=True)

    elif viz_choice == "Loadings Heatmap":
        st.markdown("### PC Loadings — What Each Component Means")

        loadings = np.array([
            [0.457, -0.083, 0.517],
            [0.452, -0.391, -0.115],
            [-0.445, 0.265, 0.702],
            [0.459, -0.079, 0.442],
            [0.422, 0.874, -0.176]
        ])
        vars_names = ["Size", "Bedrooms", "Age", "Lot", "Garage"]
        pc_names = ["PC1", "PC2", "PC3"]

        fig = go.Figure(data=go.Heatmap(
            z=loadings,
            x=pc_names, y=vars_names,
            colorscale=[[0, '#B71C1C'], [0.25, '#FFCDD2'], [0.5, 'white'],
                        [0.75, '#BBDEFB'], [1, DARK_BLUE]],
            zmid=0, zmin=-1, zmax=1,
            text=[[f"{v:.3f}" for v in row] for row in loadings],
            texttemplate="%{text}", textfont=dict(size=14),
            colorbar=dict(title="Loading")))

        fig.update_layout(
            title="PC Loadings Heatmap — Case Study (5 Variables)",
            template="plotly_white", height=450,
            font=dict(family="Times New Roman", size=13),
            yaxis=dict(autorange="reversed"))
        st.plotly_chart(fig, use_container_width=True)

        # Interpretation
        col1, col2, col3 = st.columns(3)
        with col1:
            insight_box("PC1 = Overall Bigness",
                "Similar positive loadings on Size, Beds, Lot, Garage. Negative on Age. "
                "→ Big, new houses score high.")
        with col2:
            insight_box("PC2 = Extra Garages",
                "Dominated by Garage (0.874). Other loadings small. "
                "→ Garage capacity beyond what size predicts.")
        with col3:
            insight_box("PC3 = Old but Large",
                "Age (0.702) and Size (0.517) both positive. "
                "→ Houses that are old AND large.")

    elif viz_choice == "Explained Variance Cumulative":
        st.markdown("### How Variance Accumulates Across Components")

        evals_5 = np.array([4.602, 0.228, 0.132, 0.035, 0.002])
        var_pct_5 = evals_5 / evals_5.sum() * 100
        cum_pct_5 = np.cumsum(var_pct_5)

        fig = go.Figure()
        fig.add_trace(go.Scatter(
            x=list(range(1, 6)), y=cum_pct_5,
            mode='lines+markers+text',
            line=dict(color=DARK_BLUE, width=4),
            marker=dict(size=14, color=DARK_BLUE, line=dict(color='white', width=2)),
            text=[f"{v:.1f}%" for v in cum_pct_5],
            textposition='top center', textfont=dict(size=13),
            fill='tozeroy', fillcolor='rgba(0,51,102,0.1)',
            name='Cumulative Variance'))

        # Threshold lines
        for thresh, col in [(80, GOLD), (95, 'red'), (99, 'green')]:
            fig.add_hline(y=thresh, line_dash="dash", line_color=col, opacity=0.6,
                annotation_text=f"{thresh}% threshold", annotation_position="bottom right")

        fig.update_layout(
            title="Cumulative Explained Variance — 5-Variable Case Study",
            xaxis_title="Number of Components Kept",
            yaxis_title="Cumulative Variance (%)",
            yaxis_range=[0, 105],
            xaxis=dict(tickmode='linear', dtick=1),
            template="plotly_white", height=500,
            font=dict(family="Times New Roman", size=13))
        st.plotly_chart(fig, use_container_width=True)

        gold_box("Reading This Chart",
            "Just 1 component captures 92%. Adding PC2 gets to 96.6%. "
            "With 3 of 5 components, we hit 99.2% — an excellent trade-off between "
            "simplicity and information retention.")

    footer()


# =============================================================================
# TAB 4: CASE STUDY
# =============================================================================
with tab4:
    st.markdown("## Case Study: Predicting House Prices with PCR")

    insight_box("Business Problem",
        "A real-estate analyst wants to predict house price using 5 features: "
        "Size, Bedrooms, Age, Lot, Garage. These features are highly intercorrelated. "
        "Can PCR deliver the same accuracy as MLR, without the multicollinearity headaches?")

    # Raw data
    st.markdown("### Step 1: The Raw Data")
    house_data = pd.DataFrame({
        "House": list(range(1, 13)),
        "Size (sqft)": [1500,2200,1800,2500,1200,3000,1700,2800,2000,1400,2600,1900],
        "Bedrooms": [3,4,3,4,2,5,3,4,3,2,4,3],
        "Age (yrs)": [20,5,15,8,35,3,25,10,18,30,7,22],
        "Lot (sqft)": [5000,7500,6000,8200,4200,10000,5500,9000,6500,4800,8500,6200],
        "Garage": [1,2,2,2,1,3,1,2,2,1,3,2],
        "Price ($k)": [245,380,295,425,180,540,265,475,320,210,455,305]
    })
    st.dataframe(house_data, use_container_width=True, hide_index=True)

    # Correlation matrix
    st.markdown("### Step 2: Spot the Problem — Correlation Matrix")
    features = house_data[["Size (sqft)", "Bedrooms", "Age (yrs)", "Lot (sqft)", "Garage"]]
    corr_matrix = features.corr()

    fig = go.Figure(data=go.Heatmap(
        z=corr_matrix.values,
        x=corr_matrix.columns.str.replace(r" \(.*\)", "", regex=True),
        y=corr_matrix.columns.str.replace(r" \(.*\)", "", regex=True),
        colorscale=[[0, '#B71C1C'], [0.25, '#FFCDD2'], [0.5, 'white'],
                    [0.75, '#BBDEFB'], [1, DARK_BLUE]],
        zmid=0, zmin=-1, zmax=1,
        text=[[f"{v:.3f}" for v in row] for row in corr_matrix.values],
        texttemplate="%{text}", textfont=dict(size=13)))
    fig.update_layout(
        title="Correlation Matrix — Red Flags Everywhere!",
        template="plotly_white", height=450,
        font=dict(family="Times New Roman"))
    st.plotly_chart(fig, use_container_width=True)

    warn_box("Size↔Lot: r = 0.997! Size↔Beds: r = 0.942. Beds↔Age: r = −0.946. "
             "Classic severe multicollinearity.")

    # PCA results
    st.markdown("### Step 3: Standardise & Run PCA")

    from sklearn.preprocessing import StandardScaler
    from sklearn.decomposition import PCA as skPCA
    from sklearn.linear_model import LinearRegression

    X = features.values
    y = house_data["Price ($k)"].values

    scaler = StandardScaler()
    X_std = scaler.fit_transform(X)

    pca = skPCA()
    pca.fit(X_std)
    eigenvalues = pca.explained_variance_
    var_ratio = pca.explained_variance_ratio_ * 100
    cum_var = np.cumsum(var_ratio)

    # Eigenvalue table
    ev_df = pd.DataFrame({
        "PC": [f"PC{i+1}" for i in range(5)],
        "Eigenvalue (λ)": [f"{e:.3f}" for e in eigenvalues],
        "Variance %": [f"{v:.1f}%" for v in var_ratio],
        "Cumulative %": [f"{c:.1f}%" for c in cum_var]
    })
    st.dataframe(ev_df, use_container_width=True, hide_index=True)

    gold_box("Decision",
        f"PC1 alone captures {var_ratio[0]:.1f}%. PC1+PC2+PC3 = {cum_var[2]:.1f}%. "
        "We keep k = 3 components.")

    # Scree plot for case study
    fig = make_subplots(specs=[[{"secondary_y": True}]])
    fig.add_trace(go.Bar(
        x=[f"PC{i+1}" for i in range(5)], y=var_ratio,
        marker_color=[DARK_BLUE]*3 + [LIGHT_BLUE]*2,
        text=[f"{v:.1f}%" for v in var_ratio], textposition='outside',
        name="Kept" ), secondary_y=False)
    fig.add_trace(go.Scatter(
        x=[f"PC{i+1}" for i in range(5)], y=cum_var,
        mode='lines+markers', line=dict(color='red', width=3),
        marker=dict(size=10), name="Cumulative"), secondary_y=True)
    fig.add_hline(y=95, line_dash="dash", line_color=GOLD, secondary_y=True,
        annotation_text="95% threshold")
    fig.update_yaxes(title_text="Individual %", range=[0, 100], secondary_y=False)
    fig.update_yaxes(title_text="Cumulative %", range=[0, 105], secondary_y=True)
    fig.update_layout(title="Variance Explained by Each PC", template="plotly_white",
        height=450, font=dict(family="Times New Roman"),
        legend=dict(orientation="h", y=1.08))
    st.plotly_chart(fig, use_container_width=True)

    # PCR vs MLR
    st.markdown("### Step 4: The Verdict — MLR vs PCR")

    # Fit PCR
    X_pc = pca.transform(X_std)[:, :3]
    pcr_model = LinearRegression().fit(X_pc, y)
    pcr_r2 = pcr_model.score(X_pc, y)
    pcr_pred = pcr_model.predict(X_pc)

    # Fit plain MLR
    mlr_model = LinearRegression().fit(X, y)
    mlr_r2 = mlr_model.score(X, y)
    mlr_pred = mlr_model.predict(X)

    col1, col2 = st.columns(2)
    with col1:
        st.markdown(f"""
        <div class="metric-card">
            <h2 style="color:{DARK_BLUE};">Plain MLR</h2>
            <p>R² = {mlr_r2:.4f}</p>
            <p>5 predictors | Multicollinear: YES</p>
            <p>Coefficients: Unstable, inflated SE</p>
        </div>
        """, unsafe_allow_html=True)
    with col2:
        st.markdown(f"""
        <div class="metric-card" style="border-color:{GOLD};">
            <h2 style="color:#7B6B00;">PCR (k=3) ⭐</h2>
            <p>R² = {pcr_r2:.4f}</p>
            <p>3 predictors | Multicollinear: NO</p>
            <p>Coefficients: Stable, all significant</p>
        </div>
        """, unsafe_allow_html=True)

    # Actual vs Predicted
    st.markdown("### Actual vs Predicted Price")
    fig = go.Figure()
    fig.add_trace(go.Scatter(
        x=y, y=mlr_pred, mode='markers', name='MLR',
        marker=dict(color=DARK_BLUE, size=10, symbol='circle')))
    fig.add_trace(go.Scatter(
        x=y, y=pcr_pred, mode='markers', name='PCR (k=3)',
        marker=dict(color=GOLD, size=12, symbol='diamond',
            line=dict(color=DARK_BLUE, width=1.5))))
    fig.add_trace(go.Scatter(
        x=[150, 560], y=[150, 560], mode='lines',
        line=dict(color='red', dash='dash', width=2), name='Perfect Prediction'))
    fig.update_layout(
        title="Actual vs Predicted — Both Models Perform Nearly Identically",
        xaxis_title="Actual Price ($k)", yaxis_title="Predicted Price ($k)",
        template="plotly_white", height=500,
        font=dict(family="Times New Roman"))
    st.plotly_chart(fig, use_container_width=True)

    # Comparison table
    comparison = pd.DataFrame({
        "Metric": ["R²", "Multicollinear?", "Coefficient Stability", "All Significant?",
                    "Number of Predictors", "Interpretability"],
        "Plain MLR (5 predictors)": [f"{mlr_r2:.4f}", "YES (r up to 0.997!)",
            "Unstable, inflated SE", "Some insignificant", "5", "Direct (per feature)"],
        "PCR (3 components)": [f"{pcr_r2:.4f}", "NO (PCs are orthogonal)",
            "Stable, small SE", "All |t| >> 2", "3 (reduced!)", "Abstract (use loadings)"]
    })
    st.dataframe(comparison, use_container_width=True, hide_index=True)

    gold_box("Case Study Verdict",
        "PCR delivers the <strong>same predictive accuracy</strong> as MLR but with "
        "<strong>fewer predictors, no multicollinearity, stable coefficients, and all "
        "coefficients significant</strong>. The only trade-off is interpretability.")

    footer()


# =============================================================================
# TAB 5: TRY IT YOURSELF
# =============================================================================
with tab5:
    st.markdown("## Experiment with PCA — Adjust Parameters & See Results")

    st.markdown("### Generate Your Own Correlated Dataset")

    col1, col2, col3 = st.columns(3)
    with col1:
        n_points = st.slider("Number of data points:", 20, 500, 100, 10)
    with col2:
        n_vars = st.slider("Number of variables:", 2, 10, 5)
    with col3:
        base_corr = st.slider("Base correlation strength:", 0.0, 0.99, 0.85, 0.05)

    # Generate correlated data
    np.random.seed(42)
    # Create correlation matrix
    corr = np.full((n_vars, n_vars), base_corr)
    np.fill_diagonal(corr, 1.0)
    # Cholesky decomposition for correlated data
    try:
        L = np.linalg.cholesky(corr)
        raw = np.random.randn(n_points, n_vars) @ L.T
        X_exp = raw * np.random.uniform(1, 10, n_vars) + np.random.uniform(0, 50, n_vars)

        # Standardise
        X_exp_std = (X_exp - X_exp.mean(axis=0)) / X_exp.std(axis=0, ddof=1)

        # PCA
        cov_mat = np.cov(X_exp_std, rowvar=False)
        eigenvalues_exp, eigenvectors_exp = np.linalg.eigh(cov_mat)
        idx = eigenvalues_exp.argsort()[::-1]
        eigenvalues_exp = eigenvalues_exp[idx]
        eigenvectors_exp = eigenvectors_exp[:, idx]

        var_pct_exp = eigenvalues_exp / eigenvalues_exp.sum() * 100
        cum_pct_exp = np.cumsum(var_pct_exp)

        # Display results
        st.markdown("### PCA Results for Your Dataset")

        # Metrics
        mc = st.columns(4)
        with mc[0]:
            metric_card(f"{var_pct_exp[0]:.1f}%", "PC1 Variance")
        with mc[1]:
            metric_card(f"{cum_pct_exp[min(2, n_vars-1)]:.1f}%",
                f"Top {min(3, n_vars)} PCs")
        with mc[2]:
            metric_card(f"{eigenvalues_exp[0]:.2f}", "Largest Eigenvalue")
        with mc[3]:
            metric_card(f"{eigenvalues_exp[-1]:.4f}", "Smallest Eigenvalue")

        # Choose k
        k_threshold = st.slider("Variance threshold (%) to choose k:", 70, 99, 90)
        k_chosen = np.searchsorted(cum_pct_exp, k_threshold) + 1

        st.markdown(f"""
        <div class="insight-box">
            <h4>🔍 For a {k_threshold}% variance threshold, you would keep
            <strong>k = {k_chosen}</strong> out of {n_vars} components.</h4>
        </div>
        """, unsafe_allow_html=True)

        # Scree plot
        fig = make_subplots(specs=[[{"secondary_y": True}]])
        colors = [DARK_BLUE if i < k_chosen else LIGHT_BLUE for i in range(n_vars)]
        fig.add_trace(go.Bar(
            x=[f"PC{i+1}" for i in range(n_vars)],
            y=var_pct_exp, marker_color=colors,
            text=[f"{v:.1f}%" for v in var_pct_exp], textposition='outside',
            name="Variance %"), secondary_y=False)
        fig.add_trace(go.Scatter(
            x=[f"PC{i+1}" for i in range(n_vars)],
            y=cum_pct_exp, mode='lines+markers',
            line=dict(color='red', width=3),
            marker=dict(size=8), name="Cumulative"), secondary_y=True)
        fig.add_hline(y=k_threshold, line_dash="dash", line_color=GOLD,
            secondary_y=True, annotation_text=f"{k_threshold}% threshold")
        fig.update_yaxes(title_text="Individual %", range=[0, max(var_pct_exp)*1.2], secondary_y=False)
        fig.update_yaxes(title_text="Cumulative %", range=[0, 105], secondary_y=True)
        fig.update_layout(
            title=f"Your Scree Plot — Dark bars = kept (k={k_chosen}), Light = dropped",
            template="plotly_white", height=450,
            font=dict(family="Times New Roman"),
            legend=dict(orientation="h", y=1.08))
        st.plotly_chart(fig, use_container_width=True)

        # If 2D possible, show scatter
        if n_vars >= 2:
            st.markdown("### 2D Projection (PC1 vs PC2)")
            scores = X_exp_std @ eigenvectors_exp
            fig = go.Figure()
            fig.add_trace(go.Scatter(
                x=scores[:, 0], y=scores[:, 1], mode='markers',
                marker=dict(color=DARK_BLUE, size=8, opacity=0.7,
                    line=dict(color='white', width=0.5)),
                name='PC Scores'))
            fig.add_hline(y=0, line_dash="dash", line_color="gray", opacity=0.3)
            fig.add_vline(x=0, line_dash="dash", line_color="gray", opacity=0.3)
            fig.update_layout(
                title=f"PC1 vs PC2 Scores — Correlation: {np.corrcoef(scores[:,0], scores[:,1])[0,1]:.6f}",
                xaxis_title=f"PC1 ({var_pct_exp[0]:.1f}%)",
                yaxis_title=f"PC2 ({var_pct_exp[1]:.1f}%)",
                template="plotly_white", height=500,
                font=dict(family="Times New Roman"))
            st.plotly_chart(fig, use_container_width=True)

        # Eigenvalue table
        st.markdown("### Full Eigenvalue Table")
        ev_table = pd.DataFrame({
            "PC": [f"PC{i+1}" for i in range(n_vars)],
            "Eigenvalue": [f"{e:.4f}" for e in eigenvalues_exp],
            "Variance %": [f"{v:.2f}%" for v in var_pct_exp],
            "Cumulative %": [f"{c:.2f}%" for c in cum_pct_exp],
            "Keep?": ["✅ Yes" if i < k_chosen else "❌ No" for i in range(n_vars)]
        })
        st.dataframe(ev_table, use_container_width=True, hide_index=True)

    except np.linalg.LinAlgError:
        warn_box("Could not generate data with this correlation. Try a lower value.")

    footer()


# =============================================================================
# TAB 6: LAB EXERCISE — Upload Your Own Data
# =============================================================================
with tab6:
    st.markdown("## 🔬 Lab Exercise: MLR vs PCA on Your Own Data")

    insight_box("How This Lab Works",
        "<strong>Step 1:</strong> Download the Excel template (50 rows, Y + X1–X5 with realistic correlated data).<br>"
        "<strong>Step 2:</strong> Open in Excel — modify values, add rows, change correlations, or replace with your own data.<br>"
        "<strong>Step 3:</strong> Upload the amended file here.<br>"
        "<strong>Step 4:</strong> The app runs <em>ordinary MLR</em> and then <em>PCA → PCR</em> and shows you the difference side-by-side.")

    st.markdown("---")

    # =========================================================================
    # STEP 1: GENERATE & DOWNLOAD DEFAULT TEMPLATE
    # =========================================================================
    st.markdown("### 📥 Step 1 — Download the Default Template")

    @st.cache_data
    def generate_default_template():
        """Generate a 50-row dataset with Y and 5 correlated X variables."""
        np.random.seed(2024)
        n = 50

        # Create correlated base signals
        base_size = np.random.normal(2000, 500, n)       # house size driver
        base_quality = np.random.normal(0, 1, n)          # quality driver
        base_age = np.random.normal(20, 10, n)            # age driver

        # X1: Size (sqft) — driven by base_size
        X1 = np.clip(base_size + np.random.normal(0, 80, n), 600, 5000).round(0)

        # X2: Bedrooms — highly correlated with size
        X2 = np.clip((X1 / 600 + np.random.normal(0, 0.4, n)), 1, 7).round(0)

        # X3: Age (years) — negatively correlated with size (newer houses tend to be bigger)
        X3 = np.clip(50 - X1 / 80 + np.random.normal(0, 5, n), 1, 60).round(0)

        # X4: Lot size (sqft) — very highly correlated with size
        X4 = np.clip(X1 * 3.2 + np.random.normal(0, 300, n), 2000, 16000).round(0)

        # X5: Garage spaces — correlated with size and quality
        X5 = np.clip((X1 / 1000 + base_quality * 0.3 + np.random.normal(0, 0.3, n)), 0, 4).round(0)

        # Y: Price ($000) — true relationship with noise
        Y = (
            50
            + 0.08 * X1
            + 12 * X2
            - 1.5 * X3
            + 0.01 * X4
            + 18 * X5
            + np.random.normal(0, 15, n)
        ).round(1)

        df = pd.DataFrame({
            "Y_Price_k": Y,
            "X1_Size_sqft": X1.astype(int),
            "X2_Bedrooms": X2.astype(int),
            "X3_Age_yrs": X3.astype(int),
            "X4_Lot_sqft": X4.astype(int),
            "X5_Garage": X5.astype(int),
        })
        return df

    default_df = generate_default_template()

    # Show preview
    with st.expander("👀 Preview the default template data (first 10 rows)", expanded=True):
        st.dataframe(default_df.head(10), use_container_width=True, hide_index=True)

        # Quick stats
        st.markdown("**Quick Statistics:**")
        st.dataframe(default_df.describe().round(2), use_container_width=True)

    # Create downloadable Excel
    @st.cache_data
    def create_excel_download(df):
        """Create an Excel file in memory with formatting."""
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Data', index=False)

            # Add an instructions sheet
            instructions = pd.DataFrame({
                "Instructions": [
                    "THE MOUNTAIN PATH — PCA Lab Exercise Template",
                    "",
                    "Column Descriptions:",
                    "  Y_Price_k    = House price in $000 (DEPENDENT variable)",
                    "  X1_Size_sqft = House size in square feet",
                    "  X2_Bedrooms  = Number of bedrooms",
                    "  X3_Age_yrs   = Age of house in years",
                    "  X4_Lot_sqft  = Lot size in square feet",
                    "  X5_Garage    = Number of garage spaces",
                    "",
                    "How to use:",
                    "  1. The 'Data' sheet has 50 rows of sample data.",
                    "  2. You may modify any values, add rows, or replace entirely.",
                    "  3. You may add more X columns (X6, X7, ...) — they will be auto-detected.",
                    "  4. Keep the FIRST column as Y (the response variable).",
                    "  5. All other columns are treated as X (predictor variables).",
                    "  6. Upload the modified file back to the Streamlit app.",
                    "",
                    "Tips for experimenting:",
                    "  • Make X1 and X4 nearly identical to see extreme multicollinearity.",
                    "  • Set X5 to random values to see a low-VIF variable.",
                    "  • Add 100+ rows to see how sample size affects stability.",
                    "  • Try your own real-world dataset!",
                    "",
                    "Prof. V. Ravichandran | mountainpathacademy.com",
                ]
            })
            instructions.to_excel(writer, sheet_name='Instructions', index=False, header=False)

            # Format the Data sheet
            workbook = writer.book
            ws = workbook['Data']
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            header_fill = PatternFill(start_color='003366', end_color='003366', fill_type='solid')
            header_font = Font(name='Times New Roman', bold=True, color='FFFFFF', size=11)
            cell_font = Font(name='Times New Roman', size=11)
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
                cell.border = thin_border

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.font = cell_font
                    cell.alignment = Alignment(horizontal='center')
                    cell.border = thin_border

            # Y column highlight
            y_fill = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=1):
                for cell in row:
                    cell.fill = y_fill

            # Auto-width
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 14)

            # Format Instructions sheet
            ws2 = workbook['Instructions']
            title_font = Font(name='Times New Roman', bold=True, size=14, color='003366')
            body_font = Font(name='Times New Roman', size=11)
            ws2.column_dimensions['A'].width = 70
            for idx, row in enumerate(ws2.iter_rows(min_row=1, max_row=ws2.max_row)):
                for cell in row:
                    cell.font = title_font if idx == 0 else body_font

        output.seek(0)
        return output.getvalue()

    excel_bytes = create_excel_download(default_df)

    st.download_button(
        label="⬇️  Download Excel Template (50 rows, Y + X1–X5)",
        data=excel_bytes,
        file_name="PCA_Lab_Template.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )

    st.markdown("---")

    # =========================================================================
    # STEP 2: UPLOAD FILE
    # =========================================================================
    st.markdown("### 📤 Step 2 — Upload Your (Modified) Excel File")

    uploaded_file = st.file_uploader(
        "Upload your .xlsx or .csv file",
        type=["xlsx", "csv"],
        help="First column = Y (response). All other columns = X (predictors)."
    )

    # Determine which data to use
    if uploaded_file is not None:
        if uploaded_file.name.endswith('.csv'):
            user_df = pd.read_csv(uploaded_file)
        else:
            user_df = pd.read_excel(uploaded_file, sheet_name=0)
        st.success(f"✅ Loaded **{uploaded_file.name}** — {user_df.shape[0]} rows × {user_df.shape[1]} columns")
        data_source = "uploaded"
    else:
        user_df = default_df.copy()
        data_source = "default"
        st.info("ℹ️ No file uploaded yet. Using the **default template data** for the analysis below. "
                "Upload your own file to see results on your data!")

    # Clean: drop non-numeric columns
    user_df = user_df.select_dtypes(include=[np.number]).dropna()

    if user_df.shape[1] < 3:
        st.error("❌ Need at least 1 Y column and 2 X columns (3 numeric columns minimum).")
        st.stop()

    # Split Y and X
    y_col = user_df.columns[0]
    x_cols = user_df.columns[1:]
    Y = user_df[y_col].values
    X = user_df[x_cols].values
    n_obs, p_vars = X.shape

    st.markdown(f"""
    <div class="insight-box">
        <h4>🔍 Data Summary</h4>
        <p><strong>Source:</strong> {'Uploaded file: ' + uploaded_file.name if data_source == 'uploaded' else 'Default template'}<br>
        <strong>Response (Y):</strong> {y_col} &nbsp;|&nbsp;
        <strong>Predictors:</strong> {', '.join(x_cols)} &nbsp;|&nbsp;
        <strong>Observations:</strong> {n_obs} &nbsp;|&nbsp;
        <strong>Variables:</strong> {p_vars}</p>
    </div>
    """, unsafe_allow_html=True)

    with st.expander("📋 View full dataset"):
        st.dataframe(user_df, use_container_width=True, hide_index=True)

    st.markdown("---")

    # =========================================================================
    # STEP 3: RUN ANALYSIS
    # =========================================================================
    st.markdown("### ⚙️ Step 3 — Run the Analysis")

    run_btn = st.button("🚀  Run MLR & PCA Comparison", type="primary", use_container_width=True)

    if run_btn or st.session_state.get("lab_ran", False):
        st.session_state["lab_ran"] = True

        from sklearn.preprocessing import StandardScaler
        from sklearn.decomposition import PCA as skPCA
        from sklearn.linear_model import LinearRegression

        # =================================================================
        # PART A: ORDINARY MLR
        # =================================================================
        st.markdown("---")
        st.markdown(f"""
        <div style="background:{DARK_BLUE}; color:white; padding:0.8rem 1.2rem;
                    border-radius:8px; margin-bottom:1rem;">
            <h3 style="color:white; margin:0;">Part A: Ordinary Multiple Linear Regression</h3>
        </div>
        """, unsafe_allow_html=True)

        mlr = LinearRegression().fit(X, Y)
        mlr_pred = mlr.predict(X)
        mlr_r2 = mlr.score(X, Y)
        mlr_residuals = Y - mlr_pred
        mlr_rmse = np.sqrt(np.mean(mlr_residuals**2))

        # Adjusted R²
        mlr_adj_r2 = 1 - (1 - mlr_r2) * (n_obs - 1) / (n_obs - p_vars - 1)

        # VIF calculation
        def compute_vif(X_matrix):
            vifs = []
            for j in range(X_matrix.shape[1]):
                X_others = np.delete(X_matrix, j, axis=1)
                r2_j = LinearRegression().fit(X_others, X_matrix[:, j]).score(X_others, X_matrix[:, j])
                vif_j = 1 / (1 - r2_j) if r2_j < 1 else np.inf
                vifs.append(vif_j)
            return np.array(vifs)

        vifs = compute_vif(X)

        # Coefficient table
        mlr_col1, mlr_col2 = st.columns([1.3, 1])

        with mlr_col1:
            st.markdown("#### MLR Coefficients & VIF")
            coef_df = pd.DataFrame({
                "Variable": list(x_cols),
                "Coefficient (β)": [f"{c:.4f}" for c in mlr.coef_],
                "VIF": [f"{v:.1f}" for v in vifs],
                "VIF Status": ["✅ OK" if v < 5 else ("⚠️ Moderate" if v < 10 else "🚨 SEVERE!") for v in vifs]
            })
            st.dataframe(coef_df, use_container_width=True, hide_index=True)
            st.markdown(f"**Intercept (β₀):** {mlr.intercept_:.4f}")

        with mlr_col2:
            st.markdown("#### MLR Metrics")
            metric_card(f"{mlr_r2:.4f}", "R²")
            st.markdown("")
            metric_card(f"{mlr_adj_r2:.4f}", "Adjusted R²")
            st.markdown("")
            metric_card(f"{mlr_rmse:.2f}", "RMSE")

        # Correlation heatmap
        st.markdown("#### Predictor Correlation Matrix")
        corr_mat = pd.DataFrame(X, columns=x_cols).corr()
        fig_corr = go.Figure(data=go.Heatmap(
            z=corr_mat.values, x=list(x_cols), y=list(x_cols),
            colorscale=[[0, '#B71C1C'], [0.25, '#FFCDD2'], [0.5, 'white'],
                        [0.75, '#BBDEFB'], [1, DARK_BLUE]],
            zmid=0, zmin=-1, zmax=1,
            text=[[f"{v:.3f}" for v in row] for row in corr_mat.values],
            texttemplate="%{text}", textfont=dict(size=12)))
        fig_corr.update_layout(
            title="Correlation Matrix — Look for |r| > 0.8",
            template="plotly_white", height=400,
            font=dict(family="Times New Roman"),
            yaxis=dict(autorange="reversed"))
        st.plotly_chart(fig_corr, use_container_width=True)

        # VIF bar chart
        fig_vif = go.Figure()
        vif_colors = [DARK_BLUE if v < 5 else (GOLD if v < 10 else '#B71C1C') for v in vifs]
        fig_vif.add_trace(go.Bar(
            x=list(x_cols), y=vifs, marker_color=vif_colors,
            text=[f"{v:.1f}" for v in vifs], textposition='outside'))
        fig_vif.add_hline(y=5, line_dash="dash", line_color=GOLD,
            annotation_text="Moderate (VIF=5)")
        fig_vif.add_hline(y=10, line_dash="dash", line_color="red",
            annotation_text="Severe (VIF=10)")
        fig_vif.update_layout(
            title="VIF for Each Predictor — Higher = More Multicollinearity",
            yaxis_title="VIF",
            template="plotly_white", height=400,
            font=dict(family="Times New Roman"))
        st.plotly_chart(fig_vif, use_container_width=True)

        high_vif = sum(1 for v in vifs if v > 10)
        if high_vif > 0:
            warn_box(f"<strong>{high_vif} variable(s) have VIF > 10!</strong> "
                     "This confirms severe multicollinearity. MLR coefficients are unreliable. "
                     "PCA will fix this — see Part B below.")
        elif sum(1 for v in vifs if v > 5) > 0:
            st.warning("⚠️ Some variables have moderate VIF (5–10). PCA may still improve stability.")
        else:
            gold_box("Low Multicollinearity",
                "All VIFs are below 5. MLR coefficients are reasonably stable. "
                "PCA is optional here — but let's compare anyway!")

        # =================================================================
        # PART B: PCA → PCR
        # =================================================================
        st.markdown("---")
        st.markdown(f"""
        <div style="background:linear-gradient(135deg, #7B6B00, #BFA200); color:white;
                    padding:0.8rem 1.2rem; border-radius:8px; margin-bottom:1rem;">
            <h3 style="color:white; margin:0;">Part B: PCA → Principal Component Regression</h3>
        </div>
        """, unsafe_allow_html=True)

        scaler = StandardScaler()
        X_std = scaler.fit_transform(X)

        pca = skPCA()
        pca.fit(X_std)
        eigenvalues_lab = pca.explained_variance_
        var_ratio_lab = pca.explained_variance_ratio_ * 100
        cum_var_lab = np.cumsum(var_ratio_lab)

        # Eigenvalue table
        st.markdown("#### Eigenvalues & Explained Variance")
        ev_lab_df = pd.DataFrame({
            "PC": [f"PC{i+1}" for i in range(p_vars)],
            "Eigenvalue (λ)": [f"{e:.4f}" for e in eigenvalues_lab],
            "Variance %": [f"{v:.2f}%" for v in var_ratio_lab],
            "Cumulative %": [f"{c:.2f}%" for c in cum_var_lab]
        })
        st.dataframe(ev_lab_df, use_container_width=True, hide_index=True)

        # Choose k
        st.markdown("#### Choose Number of Components (k)")
        k_slider = st.slider(
            "Variance threshold (%) — components are added until this is reached:",
            min_value=70, max_value=99, value=90, step=1,
            key="lab_k_slider"
        )
        k_chosen_lab = int(np.searchsorted(cum_var_lab, k_slider) + 1)
        k_chosen_lab = min(k_chosen_lab, p_vars)

        st.markdown(f"""
        <div class="gold-box">
            <h4>💡 Keeping k = {k_chosen_lab} component(s) out of {p_vars}
            — captures {cum_var_lab[k_chosen_lab-1]:.1f}% of total variance</h4>
        </div>
        """, unsafe_allow_html=True)

        # Scree plot
        fig_scree = make_subplots(specs=[[{"secondary_y": True}]])
        bar_colors = [DARK_BLUE if i < k_chosen_lab else LIGHT_BLUE for i in range(p_vars)]
        fig_scree.add_trace(go.Bar(
            x=[f"PC{i+1}" for i in range(p_vars)],
            y=var_ratio_lab, marker_color=bar_colors,
            text=[f"{v:.1f}%" for v in var_ratio_lab], textposition='outside',
            name="Individual %"), secondary_y=False)
        fig_scree.add_trace(go.Scatter(
            x=[f"PC{i+1}" for i in range(p_vars)],
            y=cum_var_lab, mode='lines+markers',
            line=dict(color='red', width=3), marker=dict(size=8),
            name="Cumulative %"), secondary_y=True)
        fig_scree.add_hline(y=k_slider, line_dash="dash", line_color=GOLD,
            secondary_y=True, annotation_text=f"{k_slider}% threshold")
        fig_scree.update_yaxes(title_text="Individual %",
            range=[0, max(var_ratio_lab)*1.3], secondary_y=False)
        fig_scree.update_yaxes(title_text="Cumulative %", range=[0, 105], secondary_y=True)
        fig_scree.update_layout(
            title=f"Scree Plot — Dark = kept (k={k_chosen_lab}), Light = dropped",
            template="plotly_white", height=450,
            font=dict(family="Times New Roman"),
            legend=dict(orientation="h", y=1.08))
        st.plotly_chart(fig_scree, use_container_width=True)

        # Loadings heatmap
        loadings_lab = pca.components_[:k_chosen_lab].T  # shape: (p_vars, k_chosen)
        fig_load = go.Figure(data=go.Heatmap(
            z=loadings_lab,
            x=[f"PC{i+1}" for i in range(k_chosen_lab)],
            y=list(x_cols),
            colorscale=[[0, '#B71C1C'], [0.25, '#FFCDD2'], [0.5, 'white'],
                        [0.75, '#BBDEFB'], [1, DARK_BLUE]],
            zmid=0, zmin=-1, zmax=1,
            text=[[f"{v:.3f}" for v in row] for row in loadings_lab],
            texttemplate="%{text}", textfont=dict(size=13)))
        fig_load.update_layout(
            title=f"Loadings Heatmap — How Original Variables Contribute to Each PC",
            template="plotly_white", height=max(300, p_vars * 50),
            font=dict(family="Times New Roman"),
            yaxis=dict(autorange="reversed"))
        st.plotly_chart(fig_load, use_container_width=True)

        # Fit PCR
        X_pc = pca.transform(X_std)[:, :k_chosen_lab]
        pcr = LinearRegression().fit(X_pc, Y)
        pcr_pred = pcr.predict(X_pc)
        pcr_r2 = pcr.score(X_pc, Y)
        pcr_residuals = Y - pcr_pred
        pcr_rmse = np.sqrt(np.mean(pcr_residuals**2))
        pcr_adj_r2 = 1 - (1 - pcr_r2) * (n_obs - 1) / (n_obs - k_chosen_lab - 1)

        # PCR coefficients
        st.markdown("#### PCR Coefficients")
        pcr_coef_df = pd.DataFrame({
            "Component": [f"PC{i+1}" for i in range(k_chosen_lab)],
            "Coefficient (α)": [f"{c:.4f}" for c in pcr.coef_],
            "Variance Captured": [f"{var_ratio_lab[i]:.1f}%" for i in range(k_chosen_lab)]
        })
        st.dataframe(pcr_coef_df, use_container_width=True, hide_index=True)

        # =================================================================
        # PART C: HEAD-TO-HEAD COMPARISON
        # =================================================================
        st.markdown("---")
        st.markdown(f"""
        <div style="background:linear-gradient(135deg, {DARK_BLUE}, #006644); color:white;
                    padding:0.8rem 1.2rem; border-radius:8px; margin-bottom:1rem;">
            <h3 style="color:white; margin:0;">Part C: Head-to-Head — MLR vs PCR</h3>
        </div>
        """, unsafe_allow_html=True)

        # Side-by-side metrics
        comp_col1, comp_col2 = st.columns(2)
        with comp_col1:
            st.markdown(f"""
            <div class="metric-card">
                <h2 style="color:{DARK_BLUE};">Ordinary MLR</h2>
                <hr style="border-color:{DARK_BLUE};">
                <p><strong>R²:</strong> {mlr_r2:.4f}</p>
                <p><strong>Adj R²:</strong> {mlr_adj_r2:.4f}</p>
                <p><strong>RMSE:</strong> {mlr_rmse:.2f}</p>
                <p><strong>Predictors:</strong> {p_vars} (original)</p>
                <p><strong>Max VIF:</strong> {max(vifs):.1f} {'🚨' if max(vifs)>10 else '⚠️' if max(vifs)>5 else '✅'}</p>
                <p><strong>Multicollinear?</strong> {'YES' if max(vifs)>5 else 'No'}</p>
            </div>
            """, unsafe_allow_html=True)
        with comp_col2:
            st.markdown(f"""
            <div class="metric-card" style="border-color:{GOLD};">
                <h2 style="color:#7B6B00;">PCR (k={k_chosen_lab})</h2>
                <hr style="border-color:{GOLD};">
                <p><strong>R²:</strong> {pcr_r2:.4f}</p>
                <p><strong>Adj R²:</strong> {pcr_adj_r2:.4f}</p>
                <p><strong>RMSE:</strong> {pcr_rmse:.2f}</p>
                <p><strong>Predictors:</strong> {k_chosen_lab} (components)</p>
                <p><strong>Max VIF:</strong> 1.0 ✅ (orthogonal by construction)</p>
                <p><strong>Multicollinear?</strong> Impossible (PCs are orthogonal)</p>
            </div>
            """, unsafe_allow_html=True)

        # Comparison table
        st.markdown("#### Detailed Comparison")
        comp_table = pd.DataFrame({
            "Metric": ["R²", "Adjusted R²", "RMSE", "Number of Predictors",
                       "Max VIF", "Multicollinearity?", "Coefficient Stability",
                       "Interpretability"],
            "Ordinary MLR": [
                f"{mlr_r2:.4f}", f"{mlr_adj_r2:.4f}", f"{mlr_rmse:.2f}",
                str(p_vars),
                f"{max(vifs):.1f}", "YES" if max(vifs) > 5 else "Low",
                "Unstable if VIF high" if max(vifs) > 5 else "Stable",
                "Direct (per feature)"],
            "PCR": [
                f"{pcr_r2:.4f}", f"{pcr_adj_r2:.4f}", f"{pcr_rmse:.2f}",
                str(k_chosen_lab),
                "1.0 (always)", "NO (impossible)",
                "Always stable",
                "Abstract (use loadings)"]
        })
        st.dataframe(comp_table, use_container_width=True, hide_index=True)

        # Actual vs Predicted — both models
        st.markdown("#### Actual vs Predicted — Both Models Overlaid")
        fig_avp = go.Figure()
        fig_avp.add_trace(go.Scatter(
            x=Y, y=mlr_pred, mode='markers', name=f'MLR (R²={mlr_r2:.4f})',
            marker=dict(color=DARK_BLUE, size=9, symbol='circle', opacity=0.7)))
        fig_avp.add_trace(go.Scatter(
            x=Y, y=pcr_pred, mode='markers', name=f'PCR k={k_chosen_lab} (R²={pcr_r2:.4f})',
            marker=dict(color=GOLD, size=11, symbol='diamond',
                line=dict(color=DARK_BLUE, width=1.5), opacity=0.85)))
        y_range = [min(Y)*0.9, max(Y)*1.1]
        fig_avp.add_trace(go.Scatter(
            x=y_range, y=y_range, mode='lines', name='Perfect Prediction',
            line=dict(color='red', dash='dash', width=2)))
        fig_avp.update_layout(
            title="Actual vs Predicted — How Close Are Both Models?",
            xaxis_title=f"Actual {y_col}",
            yaxis_title=f"Predicted {y_col}",
            template="plotly_white", height=500,
            font=dict(family="Times New Roman"),
            legend=dict(orientation="h", y=1.08))
        st.plotly_chart(fig_avp, use_container_width=True)

        # Residual comparison
        st.markdown("#### Residual Distribution — MLR vs PCR")
        fig_resid = make_subplots(rows=1, cols=2,
            subplot_titles=(f"MLR Residuals (RMSE={mlr_rmse:.2f})",
                            f"PCR Residuals (RMSE={pcr_rmse:.2f})"))
        fig_resid.add_trace(go.Histogram(
            x=mlr_residuals, nbinsx=15, marker_color=DARK_BLUE,
            opacity=0.7, name='MLR'), row=1, col=1)
        fig_resid.add_trace(go.Histogram(
            x=pcr_residuals, nbinsx=15, marker_color=GOLD,
            opacity=0.7, name='PCR'), row=1, col=2)
        fig_resid.update_layout(
            template="plotly_white", height=400,
            font=dict(family="Times New Roman"),
            showlegend=False)
        st.plotly_chart(fig_resid, use_container_width=True)

        # PC1 vs PC2 scatter
        if k_chosen_lab >= 2:
            st.markdown("#### PC Score Space (PC1 vs PC2)")
            all_scores = pca.transform(X_std)
            fig_pc = go.Figure()
            fig_pc.add_trace(go.Scatter(
                x=all_scores[:, 0], y=all_scores[:, 1], mode='markers',
                marker=dict(color=Y, colorscale='Viridis', size=10,
                    colorbar=dict(title=y_col), line=dict(color='white', width=0.5)),
                text=[f"{y_col}={yi:.1f}" for yi in Y],
                hovertemplate="PC1=%{x:.2f}<br>PC2=%{y:.2f}<br>%{text}"))
            fig_pc.add_hline(y=0, line_dash="dash", line_color="gray", opacity=0.3)
            fig_pc.add_vline(x=0, line_dash="dash", line_color="gray", opacity=0.3)
            fig_pc.update_layout(
                title=f"PC1 vs PC2 — Coloured by {y_col}",
                xaxis_title=f"PC1 ({var_ratio_lab[0]:.1f}%)",
                yaxis_title=f"PC2 ({var_ratio_lab[1]:.1f}%)",
                template="plotly_white", height=500,
                font=dict(family="Times New Roman"))
            st.plotly_chart(fig_pc, use_container_width=True)

        # Final verdict
        r2_diff = abs(mlr_r2 - pcr_r2)
        if max(vifs) > 10:
            verdict = (
                f"Your data has <strong>severe multicollinearity</strong> (max VIF = {max(vifs):.1f}). "
                f"PCR fixes this completely while losing only {r2_diff*100:.2f}% R². "
                f"<strong>PCR is strongly recommended</strong> over plain MLR for this dataset."
            )
        elif max(vifs) > 5:
            verdict = (
                f"Your data has <strong>moderate multicollinearity</strong> (max VIF = {max(vifs):.1f}). "
                f"PCR provides cleaner coefficient estimates. R² difference is only {r2_diff*100:.2f}%. "
                f"PCR is a <strong>good choice</strong> for improved stability."
            )
        else:
            verdict = (
                f"Your data has <strong>low multicollinearity</strong> (max VIF = {max(vifs):.1f}). "
                f"Both MLR and PCR perform similarly (R² difference = {r2_diff*100:.2f}%). "
                f"Plain MLR is fine here, but PCR still works if you want dimensionality reduction."
            )

        st.markdown(f"""
        <div class="gold-box">
            <h4>💡 Verdict for Your Data</h4>
            <p>{verdict}</p>
        </div>
        """, unsafe_allow_html=True)

    else:
        st.markdown("""
        <div style="text-align:center; padding:3rem; color:#888;">
            <h3>👆 Click "Run MLR & PCA Comparison" to start the analysis</h3>
            <p>You can run with the default data first, then upload your own file.</p>
        </div>
        """, unsafe_allow_html=True)

    footer()
